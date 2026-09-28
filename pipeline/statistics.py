"""Import pinned published inputs and construct the Phase 1 area inventory.

This module has no network access and does not import the Phase 0 audit. Source
hash verification belongs to the build orchestrator; these checks validate the
source schemas, joins, raw values and the calculation contract independently.
"""

from __future__ import annotations

import csv
import json
import re
from collections import Counter
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from fractions import Fraction
from io import TextIOWrapper
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import openpyxl

BANDS = "ABCDEFGH"
NINTHS = (6, 7, 8, 9, 11, 13, 15, 18)
PRICE_PERIOD = "Year ending Sep 2025"
STOCK_MEMBER = "CTSOP1.1/CTSOP1_1_2025_03_31.csv"
STOCK_COLUMNS = ["geography", "ba_code", "ecode", "area_name"] + [
    "band_" + band.lower() for band in BANDS + "I"
] + ["all_properties"]
MSOA_CODE = re.compile(r"[EW]020\d{5}\Z")
LAD_CODE = re.compile(r"(?:E0[6789]0|W060)\d{5}\Z")
KNOWN_OLD_LADS = {"E08000016": "E08000038", "E08000019": "E08000039"}
KNOWN_PARENT_DIFFERENCES = {
    ("E08000016", "E08000038"): 30,
    ("E08000019", "E08000039"): 70,
    ("E07000082", "E07000081"): 1,
}
NON_NESTING_LADS = {"E07000081", "E07000082"}  # Gloucester / Stroud, MSOA E02004651.


def require(condition: bool, message: str) -> None:
    if not condition:
        raise ValueError(message)


def unique_insert(records: dict, code: str, value: dict, source: str) -> None:
    require(code not in records, f"Duplicate {source} code: {code}")
    records[code] = value


def money_pence(raw: object) -> int | None:
    """Convert a finite numeric source currency to integer pence, half up.

    A missing or symbolic cell is not a number. The caller classifies its raw
    marker, which remains in sourceRefs even when the estimate is unavailable.
    """
    if raw is None or isinstance(raw, bool):
        return None
    try:
        value = Decimal(str(raw).strip())
        if not value.is_finite():
            return None
        return int((value * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    except (InvalidOperation, ValueError, OverflowError):
        return None


def read_prices(path: Path, geography: str) -> dict[str, dict]:
    sheet = "1a" if geography == "MSOA" else "2a"
    expected_code_header = "MSOA code" if geography == "MSOA" else "Local authority code"
    matcher = MSOA_CODE if geography == "MSOA" else LAD_CODE
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        require(sheet in workbook.sheetnames, f"Missing price sheet {sheet}")
        rows = workbook[sheet].iter_rows(values_only=True)
        next(rows)
        next(rows)
        headers = next(rows)
        require(len(headers) == 124 and headers[2] == expected_code_header
                and headers[123] == PRICE_PERIOD,
                f"Unexpected {geography} price schema or period")
        require(headers[0] == ("Local authority code" if geography == "MSOA"
                               else "Region/Country code"), "Unexpected price parent column")
        result = {}
        for row_number, row in enumerate(rows, 4):
            code = row[2]
            if not isinstance(code, str) or not code.startswith(("E0", "W0")):
                continue
            require(bool(matcher.fullmatch(code)), f"Invalid {geography} price code: {code}")
            unique_insert(result, code, {
                "sourceId": "ons-msoa-prices" if geography == "MSOA" else "ons-lad-prices",
                "code": code, "name": row[3], "rawValue": row[123],
                "cell": f"{sheet}!DT{row_number}", "codeCell": f"{sheet}!C{row_number}",
                "parentCode": row[0] if geography == "MSOA" else None,
                "period": "2024-10-01/2025-09-30", "units": "GBP",
            }, geography + " price")
        return result
    finally:
        workbook.close()


def read_stock(path: Path) -> tuple[dict[str, dict], dict[str, dict], dict]:
    """Import separate MSOA and LAUA keys; never mix in LSOA/CTYMET totals."""
    msoas, lads, unmatched = {}, {}, {}
    parent = None
    geography_counts = Counter()
    with ZipFile(path) as archive, archive.open(STOCK_MEMBER) as stream:
        rows = csv.DictReader(TextIOWrapper(stream, encoding="utf-8-sig"))
        require(rows.fieldnames == STOCK_COLUMNS, "Unexpected CTSOP1.1 stock columns")
        for row_number, row in enumerate(rows, 2):
            require(None not in row and all(value is not None for value in row.values()),
                    f"Malformed stock CSV row {row_number}")
            geography_counts[row["geography"]] += 1
            if row["geography"] == "LAUA":
                parent = row["ecode"]
            if row["geography"] not in ("MSOA", "LAUA", "UNMD"):
                continue
            code = row["ecode"]
            if row["geography"] != "UNMD":
                matcher = MSOA_CODE if row["geography"] == "MSOA" else LAD_CODE
                require(bool(matcher.fullmatch(code)), f"Invalid stock code {code}")
                require(not code.startswith("E") or row["band_i"] == "..",
                        f"Unexpected English band I: {code}")
            record = {
                "sourceId": "voa-stock-all", "member": STOCK_MEMBER,
                "sourceRow": row_number, "code": code, "name": row["area_name"],
                "geography": row["geography"], "billingAuthorityCode": row["ba_code"],
                "rawCounts": {b: row["band_" + b.lower()] for b in BANDS},
                "rawReportedTotal": row["all_properties"], "rawBandI": row["band_i"],
                "observedParentGroup": parent if row["geography"] == "MSOA" else None,
                "period": "2025-03-31", "geographyAssignment": "NSPL February 2025 postcodes",
            }
            target = msoas if row["geography"] == "MSOA" else lads if row["geography"] == "LAUA" else unmatched
            unique_insert(target, code, record, row["geography"] + " stock")
    return msoas, lads, {"geographyRowCounts": dict(sorted(geography_counts.items())),
                        "unmatchedStockRows": list(unmatched.values())}


def read_billing(path: Path) -> dict[str, dict]:
    namespaces = {"table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
                  "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0"}
    table_ns, office_ns = ("{" + namespaces[key] + "}" for key in ("table", "office"))
    with ZipFile(path) as archive:
        tree = ET.fromstring(archive.read("content.xml"))
    tables = [table for table in tree.findall(".//table:table", namespaces)
              if table.get(table_ns + "name") == "Data_Billing"]
    require(len(tables) == 1, "Missing or repeated Data_Billing sheet")
    result, row_number, header_checked = {}, 1, False
    for row in tables[0].findall("table:table-row", namespaces):
        values = []
        repetitions = int(row.get(table_ns + "number-rows-repeated", "1"))
        for cell in row:
            value = (cell.get(office_ns + "value")
                     if cell.get(office_ns + "value-type") in ("float", "currency", "percentage")
                     else "".join(cell.itertext()))
            copies = int(cell.get(table_ns + "number-columns-repeated", "1"))
            values.extend([value] * min(copies, max(0, 38 - len(values))))
        if row_number == 5:
            require(len(values) == 38 and values[1] == "ONS Code" and values[2] == "Authority"
                    and "including both local and major precepts" in values[37],
                    "Unexpected billing columns B/C/AL")
            header_checked = True
        if len(values) >= 38 and values[1].startswith("E0"):
            code = values[1]
            require(bool(LAD_CODE.fullmatch(code)) and repetitions == 1,
                    f"Invalid or repeated billing row {row_number}")
            unique_insert(result, code, {
                "sourceId": "council-tax", "code": code, "name": values[2],
                "codeCell": f"Data_Billing!B{row_number}",
                "cell": f"Data_Billing!AL{row_number}", "rawBandD": values[37],
                "bandDPence": money_pence(values[37]), "period": "2026-04-01/2027-03-31",
                "rawComponents": dict(zip(("Z", "AH", "AI", "AJ", "AK"),
                                           (values[25], *values[33:37]))),
            }, "billing")
        row_number += repetitions
    require(header_checked, "Billing header absent")
    return result


def read_lookup(path: Path) -> dict[str, dict]:
    result = {}
    with path.open(encoding="utf-8-sig", newline="") as stream:
        rows = csv.DictReader(stream)
        require({"MSOA21CD", "MSOA21NM", "LAD25CD", "LAD25NM"}.issubset(rows.fieldnames or []),
                "Unexpected MSOA/LAD lookup schema")
        for row_number, row in enumerate(rows, 2):
            code, parent = row["MSOA21CD"], row["LAD25CD"]
            require(bool(MSOA_CODE.fullmatch(code)) and bool(LAD_CODE.fullmatch(parent)),
                    f"Invalid lookup row {row_number}")
            unique_insert(result, code, {"sourceId": "msoa-lad-lookup", "sourceRow": row_number,
                                        "code": code, "name": row["MSOA21NM"],
                                        "parentCode": parent, "parentName": row["LAD25NM"]}, "lookup")
    return result


def boundary_inventory(path: Path, field: str) -> dict[str, str]:
    collection = json.loads(path.read_text())
    require(collection.get("type") == "FeatureCollection" and not collection.get("exceededTransferLimit"),
            f"Invalid or truncated boundary collection: {path.name}")
    result = {}
    for feature in collection["features"]:
        code = feature["properties"][field]
        unique_insert(result, code, feature["properties"].get(field.replace("CD", "NM"), code), "boundary")
    return result


def stock_values(record: dict | None) -> tuple[dict, list[str]]:
    if record is None:
        return {}, ["stock-unavailable"]
    reasons, counts = [], {}
    for band in BANDS:
        raw = record["rawCounts"].get(band)
        if isinstance(raw, str) and re.fullmatch(r"[0-9]+", raw):
            counts[band] = int(raw)
        else:
            reasons.append("stock-marker-unverified" if raw == "-" else "stock-suppressed" if raw == "[c]"
                           else "stock-unavailable" if raw in (None, "") else "source-value-unrecognised")
    raw_total = record.get("rawReportedTotal")
    total = int(raw_total) if isinstance(raw_total, str) and re.fullmatch(r"[0-9]+", raw_total) else None
    if total is None:
        reasons.append("stock-total-unavailable")
    result = {"counts": counts, "reportedTotal": total, "sumOfBands": None,
              "weightedNinthsSum": None, "reconciliationDifference": None}
    if len(counts) == 8:
        summed = sum(counts.values())
        result.update(sumOfBands=summed, weightedNinthsSum=sum(counts[b] * n for b, n in zip(BANDS, NINTHS)))
        if not summed:
            reasons.append("stock-denominator-zero")
        if total is not None:
            difference = summed - total
            require(abs(difference) <= 45, f"Stock rounding bound exceeded: {record['code']} ({difference})")
            result["reconciliationDifference"] = difference
    return result, sorted(set(reasons))


def make_area(code: str, name: str, geography: str, parent: str | None,
              price: dict | None, stock: dict | None, charge: dict | None,
              lookup: dict | None = None, geography_review: bool = False) -> dict:
    reasons, flags = [], ["rounded-stock-counts", "authority-average-charge-proxy", "mixed-source-periods"]
    raw_price = price["rawValue"] if price else None
    price_pence = money_pence(raw_price)
    if price_pence is None or price_pence <= 0:
        price_pence = None
        reasons.append("price-unavailable")
        if raw_price not in (None, "", "[x]", "[c]", "..", "-", 0) and money_pence(raw_price) is None:
            reasons.append("source-value-unrecognised")
    values, stock_reasons = stock_values(stock)
    reasons.extend(stock_reasons)
    band_d = charge["bandDPence"] if charge else None
    if band_d is None or band_d <= 0:
        reasons.append("charge-unavailable")
    if charge:
        components = [money_pence(value) for value in charge["rawComponents"].values()]
        if band_d is not None and all(value is not None for value in components):
            require(abs(sum(components) - band_d) <= 5, f"Charge components fail reconciliation: {code}")
    if geography_review:
        reasons.append("geography-needs-review")
        flags.append("geography-needs-review")
    if parent in KNOWN_OLD_LADS.values():
        flags.append("authority-boundary-change")
    if price_pence is not None and not 1_000_000 <= price_pence <= 1_000_000_000:
        flags.append("price-outlier-review")
    if band_d is not None and not 10_000 <= band_d <= 1_000_000:
        flags.append("charge-outlier-review")
    baseline = None
    if not stock_reasons and band_d is not None and band_d > 0 and not geography_review:
        exact = Fraction(band_d * values["weightedNinthsSum"], 9 * values["sumOfBands"])
        baseline = {"numerator": exact.numerator, "denominator": exact.denominator}
    return {
        "code": code, "name": name, "geography": geography, "parentCode": parent,
        "propertyType": "all", "availability": "unavailable" if reasons else "available",
        "unavailableReasons": sorted(set(reasons)), "qualityFlags": sorted(set(flags)),
        "pricePence": price_pence, "councilTaxExactPence": baseline, "transactionCount": None,
        "sourceRefs": {"price": price, "stock": {**stock, **values} if stock else None,
                       "charge": charge, "parentLookup": lookup},
    }


def validate_parent_groups(stock: dict[str, dict], lookup: dict[str, dict]) -> list[dict]:
    differences = []
    for code in sorted(stock.keys() & lookup.keys()):
        if not code.startswith("E"):
            continue
        old, current = stock[code]["observedParentGroup"], lookup[code]["parentCode"]
        if old == current:
            continue
        require((old, current) in KNOWN_PARENT_DIFFERENCES, f"Unexplained stock parent join: {code}")
        require((old, current) != ("E07000082", "E07000081") or code == "E02004651",
                f"Unexpected Gloucester/Stroud stock parent join: {code}")
        differences.append({"code": code, "stockParentGroup": old, "currentParentCode": current,
                            "reason": "2025-authority-boundary-change" if old in KNOWN_OLD_LADS
                            else "known-non-nesting-msoa"})
    require(Counter((d["stockParentGroup"], d["currentParentCode"]) for d in differences)
            == KNOWN_PARENT_DIFFERENCES, "Reviewed stock-parent disagreement inventory changed")
    return differences


def validate_lad_example(root: Path, areas: list[dict]) -> dict:
    """Check independently transcribed expectations, never rewrite the fixture."""
    fixture = json.loads((root / "data/fixtures/phase1-lad-worked-example.json").read_text())
    area = next(area for area in areas if area["code"] == fixture["code"])
    require(area["pricePence"] == fixture["price"]["pricePence"]
            and area["councilTaxExactPence"] == fixture["expected"]["councilTaxExactPence"],
            "Independent LAD price/baseline fixture does not reconcile")
    for source, fields in (("price", ("sourceId", "cell", "rawValue")),
                           ("stock", ("sourceId", "member", "sourceRow", "geography", "billingAuthorityCode",
                                      "rawCounts", "rawReportedTotal", "sumOfBands", "weightedNinthsSum")),
                           ("charge", ("sourceId", "cell", "rawBandD", "bandDPence"))):
        require(all(area["sourceRefs"][source][key] == fixture[source][key] for key in fields),
                f"Independent LAD {source} fixture does not reconcile")
    return {"code": fixture["code"], "fixture": "data/fixtures/phase1-lad-worked-example.json",
            "inputsMatch": True, "exactBaselineMatches": True}


def select_geographies(root: Path, english: set[str], english_lads: set[str], lookup: dict,
                       scope: str = "sample") -> list[dict]:
    """Choose the complete requested inventory, never just its usable records."""
    require(scope in {"sample", "england"}, f"Unsupported geographic scope: {scope}")
    if scope == "sample":
        selections = json.loads((root / "data/fixtures/phase0-sample-areas.json").read_text())
        lads = {sample["lad_code"] for sample in selections}
        msoas = {code for sample in selections for code in sample["msoa_codes"]}
        require(len(lads) == len(selections) == 5 and len(msoas) == 198,
                "Sample selection inventory changed")
    else:
        selections = [{"lad_code": lad,
                       "msoa_codes": sorted(code for code in english if lookup[code]["parentCode"] == lad)}
                      for lad in sorted(english_lads)]
        for item in selections:
            item["msoa_count"] = len(item["msoa_codes"])
    for selection in selections:
        selected = {code for code in english if lookup[code]["parentCode"] == selection["lad_code"]}
        require(selected == set(selection["msoa_codes"]) and len(selected) == selection["msoa_count"],
                f"Selected parent lookup changed: {selection['lad_code']}")
    return selections


def reconcile_lad_stock(selections: list[dict], msoa_stock: dict, lad_stock: dict,
                        scope: str = "sample") -> tuple[list[dict], list[dict], list[dict]]:
    """Cross-check rounded totals without pretending best-fit MSOAs nest exactly.

    Gloucester and Stroud keep independent published LAD inputs. Their reviewed
    best-fit difference is audited against the stock source's own row grouping;
    no MSOA count or old LAD total is reassigned to manufacture agreement.
    """
    reconciliations, non_nesting, exclusions = [], [], []
    for selection in selections:
        code, members = selection["lad_code"], selection["msoa_codes"]
        if code not in lad_stock:
            require(scope == "england" and code in KNOWN_OLD_LADS.values(),
                    f"Missing LAD stock for reconciliation: {code}")
            exclusions.append({"code": code, "reason": "current-lad-stock-vintage-unavailable"})
            continue
        raw_totals = [msoa_stock[c]["rawReportedTotal"] for c in members if c in msoa_stock]
        parent_total = lad_stock[code]["rawReportedTotal"]
        if len(raw_totals) != len(members) or not all(re.fullmatch(r"[0-9]+", v) for v in raw_totals + [parent_total]):
            if scope == "england":
                exclusions.append({"code": code, "reason": "published-total-unavailable"})
            continue
        total, reported = sum(map(int, raw_totals)), int(parent_total)
        tolerance = 5 * (len(members) + 1)
        detail = {"code": code, "summedMsoaPublishedTotals": total,
                  "ladPublishedTotal": reported, "difference": total - reported,
                  "tolerance": tolerance, "msoaCount": len(members)}
        if abs(total - reported) > tolerance:
            require(scope == "england" and code in NON_NESTING_LADS,
                    f"MSOA/LAD stock total reconciliation failed: {code}")
            native = sorted(c for c, row in msoa_stock.items() if row["observedParentGroup"] == code)
            require(set(native) ^ set(members) == {"E02004651"},
                    f"Unexplained non-nesting stock inventory: {code}")
            native_raw = [msoa_stock[c]["rawReportedTotal"] for c in native]
            require(all(re.fullmatch(r"[0-9]+", v) for v in native_raw),
                    f"Unresolved native stock totals: {code}")
            native_sum, native_tolerance = sum(map(int, native_raw)), 5 * (len(native) + 1)
            require(abs(native_sum - reported) <= native_tolerance,
                    f"Native stock-group total reconciliation failed: {code}")
            non_nesting.append({**detail, "reason": "reviewed-best-fit-msoa-does-not-nest",
                                "msoaCode": "E02004651", "stockGroupedMsoaTotal": native_sum,
                                "stockGroupedDifference": native_sum - reported,
                                "stockGroupedTolerance": native_tolerance})
        else:
            reconciliations.append(detail)
    return reconciliations, non_nesting, exclusions


def build_statistics(root: Path, sources: dict, geography_review: set[str] | None = None,
                     *, scope: str = "sample"
                     ) -> tuple[list[dict], dict, dict]:
    source_path = lambda key: root / sources[key]["localPath"]
    msoa_prices = read_prices(source_path("ons-msoa-prices"), "MSOA")
    lad_prices = read_prices(source_path("ons-lad-prices"), "LAD")
    msoa_stock, lad_stock, stock_report = read_stock(source_path("voa-stock-all"))
    billing = read_billing(source_path("council-tax"))
    lookup = read_lookup(source_path("msoa-lad-lookup"))
    msoa_boundaries = boundary_inventory(source_path("msoa-boundaries"), "MSOA21CD")
    lad_boundaries = boundary_inventory(source_path("lad-boundaries"), "LAD25CD")
    metadata = json.loads(source_path("lad-boundary-metadata").read_text())
    require(metadata["name"] == "LAD_MAY_2025_UK_BSC_V2", "LAD boundary vintage changed")
    require(set(lookup) == set(msoa_boundaries) and len(lookup) == 7264, "MSOA lookup/boundary inventories differ")
    english = {code for code in lookup if code.startswith("E")}
    english_lads = {code for code in lad_boundaries if code.startswith("E")}
    require(len(english) == 6856 and len(english_lads) == 296, "Pinned England inventory changed")
    for label, actual, expected in (("MSOA price", msoa_prices, lookup), ("MSOA stock", msoa_stock, lookup),
                                    ("billing", billing, english_lads)):
        require(not set(actual) - set(expected), f"Unexpected {label} codes: {sorted(set(actual) - set(expected))}")
    require(all(value["parentCode"] == lookup[code]["parentCode"] for code, value in msoa_prices.items()),
            "ONS price parent conflicts with official lookup")
    require(all(lookup[code]["parentCode"] in english_lads for code in english), "Unmatched official LAD parent")
    stock_lads = {code for code in lad_stock if code.startswith("E")}
    require(stock_lads - english_lads == set(KNOWN_OLD_LADS)
            and english_lads - stock_lads == set(KNOWN_OLD_LADS.values()),
            "Unexplained LAD stock/boundary vintage mismatch")
    require({code for code in lad_prices if code.startswith("E")} <= english_lads,
            "Unexpected English LAD price code")
    parent_differences = validate_parent_groups(msoa_stock, lookup)
    selections = select_geographies(root, english, english_lads, lookup, scope)
    sample_lads = {sample["lad_code"] for sample in selections}
    sample_msoas = {code for sample in selections for code in sample["msoa_codes"]}
    review = set(geography_review or ()) | {"E02004651"}
    areas = []
    for code in sorted(sample_msoas):
        parent = lookup[code]["parentCode"]
        areas.append(make_area(code, msoa_prices.get(code, lookup[code])["name"], "MSOA", parent,
                               msoa_prices.get(code), msoa_stock.get(code), billing.get(parent),
                               lookup[code], code in review))
    for code in sorted(sample_lads):
        areas.append(make_area(code, lad_boundaries[code], "LAD", None, lad_prices.get(code),
                               lad_stock.get(code), billing.get(code), geography_review=code in review))
    reconciliations, non_nesting, reconciliation_exclusions = reconcile_lad_stock(
        selections, msoa_stock, lad_stock, scope)
    national_markers = Counter(raw for code in sorted(english & msoa_stock.keys())
                              for raw in msoa_stock[code]["rawCounts"].values()
                              if not re.fullmatch(r"[0-9]+", raw))
    national_reconciliation = []
    for code in sorted(english & msoa_stock.keys()):
        values, _ = stock_values(msoa_stock[code])
        if values.get("reconciliationDifference") is not None:
            national_reconciliation.append(values["reconciliationDifference"])
    report = {
        "duplicates": [], "unexplainedJoins": [],
        "independentLadExample": validate_lad_example(root, areas),
        "inputInventory": {"englishMsoas": len(english), "englishLads": len(english_lads),
                           "msoaPrices": len(msoa_prices), "msoaStock": len(msoa_stock),
                           "ladPrices": len(lad_prices), "englishLadStock": len(stock_lads), "charges": len(billing)},
        "missingInputs": {"msoaPrices": sorted(set(lookup) - msoa_prices.keys()),
                          "msoaStock": sorted(set(lookup) - msoa_stock.keys()),
                          "ladPrices": sorted(english_lads - lad_prices.keys()),
                          "charges": sorted(english_lads - billing.keys())},
        "stockParentDisagreements": parent_differences,
        "ladVintageExclusions": [{"stockCode": old, "currentCode": new,
                                  "reason": "2025-boundary-change; old totals must not be relabelled"}
                                 for old, new in sorted(KNOWN_OLD_LADS.items())],
        "stock": {**stock_report, "nationalMsoaNonnumericBandMarkers": dict(sorted(national_markers.items())),
                  "roundingToleranceDwellings": 45,
                  "maxAbsoluteMsoaBandReconciliationDifference": max(map(abs, national_reconciliation), default=0),
                  ("sampleLadPublishedTotalReconciliations" if scope == "sample"
                   else "nationalLadPublishedTotalReconciliations"): reconciliations,
                  **({"reviewedNonNestingReconciliations": non_nesting,
                      "ladReconciliationExclusions": reconciliation_exclusions} if scope == "england" else {})},
        "coverage": {geography: {"total": sum(a["geography"] == geography for a in areas),
                                  "available": sum(a["geography"] == geography and a["availability"] == "available" for a in areas),
                                  "unavailable": sum(a["geography"] == geography and a["availability"] == "unavailable" for a in areas)}
                     for geography in ("MSOA", "LAD")},
        "unavailableReasonCounts": dict(sorted(Counter(reason for area in areas for reason in area["unavailableReasons"]).items())),
        "outliers": [{"code": area["code"], "flags": [f for f in area["qualityFlags"] if "outlier" in f]}
                     for area in areas if any("outlier" in f for f in area["qualityFlags"])],
        "sampleGeographyReview": sorted(review & (sample_msoas | sample_lads)),
    }
    return areas, report, {"parent_lookup": {code: value["parentCode"] for code, value in lookup.items()},
                           "sample_msoas": sample_msoas, "sample_lads": sample_lads}
